"""Parser for the TEYDEB programme limits workbook (one TABLO section per sheet).

Tables are turned into sentences a retriever can match. In a sheet whose header row lists
programme codes (``Limitler``: rows = criteria, columns = programmes) each row becomes
"Proje Destek Süresi — 1501: 36 AY; 1507: 18 AY", so a value never loses its programme.
Other sheets are written row by row as "cell | cell".
"""

import re
from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook

from tubitak_rag.domain.document import Document, Paragraph, Section
from tubitak_rag.domain.enums import SectionKind
from tubitak_rag.domain.source import SourceEntry
from tubitak_rag.ingestion.document_builder import build_document_skeleton, make_section
from tubitak_rag.ingestion.hashing import sha256_of
from tubitak_rag.ingestion.normalize import collapse_ws, normalize_text

PROGRAM_CODE_CELL_RE = re.compile(r"^\d{4}$")
_MIN_CODE_COLUMNS = 2


def cell_text(value: object) -> str:
    """Cell value as one clean line (integral floats lose their ".0")."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return collapse_ws(normalize_text(str(value)))


def _code_columns(row: Sequence[str]) -> dict[int, str] | None:
    codes = {i: cell for i, cell in enumerate(row) if PROGRAM_CODE_CELL_RE.match(cell)}
    return codes if len(codes) >= _MIN_CODE_COLUMNS else None


def sheet_lines(rows: Sequence[Sequence[str]]) -> list[str]:
    """Text lines of one sheet (empty rows are skipped)."""
    lines: list[str] = []
    codes: dict[int, str] | None = None
    label = ""
    for row in rows:
        if not any(row):
            continue
        if codes is None and (codes := _code_columns(row)):
            continue
        if codes is None:
            lines.append(" | ".join(cell for cell in row if cell))
            continue
        label = row[0] or label
        values = "; ".join(f"{codes[i]}: {row[i]}" for i in codes if i < len(row) and row[i])
        lines.append(f"{label} — {values}" if label and values else values or label)
    return lines


class XlsxParser:
    """``DocumentParser`` for XLSX tables."""

    def parse(self, entry: SourceEntry, path: Path) -> Document:
        """One TABLO section per non-empty sheet."""
        workbook = load_workbook(path, read_only=True, data_only=True)
        sections: list[Section] = []
        try:
            for sheet in workbook.worksheets:
                rows = [[cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
                lines = sheet_lines(rows)
                if not lines:
                    continue
                sections.append(
                    make_section(
                        entry.doc_id,
                        len(sections) + 1,
                        SectionKind.TABLO,
                        text="\n".join(lines),
                        paragraphs=[Paragraph(label=None, text=ln, page=None) for ln in lines],
                        title=sheet.title,
                    )
                )
        finally:
            workbook.close()
        doc = build_document_skeleton(entry, sha256_of(path))
        return doc.model_copy(update={"sections": sections})
