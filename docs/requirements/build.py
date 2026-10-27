"""Build the requirements workbook and Markdown table from ihtiyaclar.yaml.

Usage:
    python docs/requirements/build.py

Outputs:
    docs/ihtiyaclar.xlsx            professional tracking workbook
    docs/ihtiyaclar.md              table section between IHTIYACLAR markers is regenerated
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "requirements" / "ihtiyaclar.yaml"
XLSX_OUT = ROOT / "docs" / "ihtiyaclar.xlsx"
MD_DOC = ROOT / "docs" / "ihtiyaclar.md"
MD_START, MD_END = "<!-- IHTIYACLAR:BASLA -->", "<!-- IHTIYACLAR:BITIS -->"

STATUSES = ["Bekliyor", "Karşılandı", "Kısmi", "Karşılanmadı", "Kapsam dışı"]
PRIORITIES = ["Zorunlu", "Önemli", "İsteğe bağlı"]
METHODS = ["Test", "Gösterim", "İnceleme", "Analiz"]
SENTENCE_PATTERN = re.compile(r"(malıdır|melidir)\.$")

FONT = "Arial"
NAVY = "1F3A5F"
GREY = "F2F2F2"
LINE = "BFBFBF"
STATUS_FILLS = {
    "Karşılandı": "D9EAD3",
    "Kısmi": "FFF2CC",
    "Karşılanmadı": "F4CCCC",
    "Kapsam dışı": "E7E6E6",
}

COLUMNS = [
    ("ID", 11),
    ("Kategori", 22),
    ("EARS Tipi", 12),
    ("İhtiyaç", 70),
    ("Kabul Kriteri", 46),
    ("Öncelik", 11),
    ("Doğrulama", 12),
    ("Durum", 14),
    ("✓", 5),
    ("Kanıt", 40),
    ("Not", 30),
]

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def load() -> dict:
    data = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))
    ids = [r["id"] for r in data["ihtiyaclar"]]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate requirement IDs")
    for req in data["ihtiyaclar"]:
        if not SENTENCE_PATTERN.search(req["ihtiyac"]):
            raise ValueError(f"{req['id']}: sentence must end with -malıdır/-melidir.")
        req.setdefault("durum", "Bekliyor")
        req.setdefault("kanit", "")
        req.setdefault("not", "")
        for field, allowed in (("durum", STATUSES), ("oncelik", PRIORITIES), ("dogrulama", METHODS)):
            if req[field] not in allowed:
                raise ValueError(f"{req['id']}: invalid {field} '{req[field]}'")
    return data


def font(**kw) -> Font:
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def header_row(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def build_requirements_sheet(wb: Workbook, data: dict) -> int:
    ws = wb.create_sheet("İhtiyaçlar")
    header_row(ws, 1, [c for c, _ in COLUMNS])
    for idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 28

    cats = data["kategoriler"]
    for r, req in enumerate(data["ihtiyaclar"], start=2):
        values = [
            req["id"],
            cats[req["kat"]],
            req["tip"],
            req["ihtiyac"],
            req["kabul"],
            req["oncelik"],
            req["dogrulama"],
            req["durum"],
            f'=IF(H{r}="Karşılandı","✔",IF(H{r}="Kısmi","◐",IF(H{r}="Karşılanmadı","✖","")))',
            req["kanit"],
            req["not"],
        ]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = font(bold=(c == 1))
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=c in (2, 4, 5, 10, 11),
                horizontal="center" if c in (3, 6, 7, 8, 9) else "left",
            )
    last = len(data["ihtiyaclar"]) + 1

    status_dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    status_dv.error, status_dv.errorTitle = "Listeden bir durum seçin.", "Geçersiz durum"
    prio_dv = DataValidation(type="list", formula1='"' + ",".join(PRIORITIES) + '"')
    ws.add_data_validation(status_dv)
    ws.add_data_validation(prio_dv)
    status_dv.add(f"H2:H{last}")
    prio_dv.add(f"F2:F{last}")

    for status, color in STATUS_FILLS.items():
        ws.conditional_formatting.add(
            f"A2:K{last}",
            FormulaRule(formula=[f'$H2="{status}"'], fill=PatternFill("solid", fgColor=color)),
        )

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:K{last}"
    ws.sheet_view.zoomScale = 90
    return last


def build_summary_sheet(wb: Workbook, data: dict, last: int) -> None:
    ws = wb.active
    ws.title = "Özet"
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDEFGH", (34, 13, 13, 13, 13, 13, 13, 13)):
        ws.column_dimensions[col].width = width

    meta = data["meta"]
    ws["A1"] = f"{meta['proje']} — İhtiyaç Karşılama Matrisi"
    ws["A1"].font = font(size=16, bold=True, color=NAVY)
    ws["A2"] = f"Sürüm {meta['surum']} · Sahip: {meta['sahip']} · Oluşturma: {date.today().isoformat()}"
    ws["A2"].font = font(color="595959")
    ws["A3"] = (
        "Kaynak: docs/requirements/ihtiyaclar.yaml (tek kaynak). Bu dosya build.py ile üretilir; "
        "durum güncellemeleri YAML'da yapılıp dosya yeniden üretilmelidir."
    )
    ws["A3"].font = font(italic=True, color="595959", size=9)

    req = "'İhtiyaçlar'"
    rng_status = f"{req}!$H$2:$H${last}"
    rng_cat = f"{req}!$B$2:$B${last}"
    rng_prio = f"{req}!$F$2:$F${last}"

    ws["A5"] = "Genel durum"
    ws["A5"].font = font(size=12, bold=True, color=NAVY)
    header_row(ws, 6, ["Gösterge", "Değer"])
    rows = [("Toplam ihtiyaç", f"=COUNTA({req}!$A$2:$A${last})")]
    rows += [(s, f'=COUNTIF({rng_status},"{s}")') for s in STATUSES]
    rows.append(("Karşılanma oranı (kapsam içi)", "=IFERROR(B9/(B7-B12),0)"))
    rows.append(
        ("Zorunlu ihtiyaçlarda karşılanma", f'=IFERROR(COUNTIFS({rng_prio},"Zorunlu",{rng_status},"Karşılandı")/COUNTIF({rng_prio},"Zorunlu"),0)')
    )
    for i, (label, formula) in enumerate(rows, start=7):
        ws.cell(row=i, column=1, value=label).font = font()
        cell = ws.cell(row=i, column=2, value=formula)
        cell.font = font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        if "oranı" in label or "karşılanma" in label:
            cell.number_format = "0.0%"
        for c in (1, 2):
            ws.cell(row=i, column=c).border = BORDER

    start = 7 + len(rows) + 2
    ws.cell(row=start - 1, column=1, value="Kategoriye göre").font = font(size=12, bold=True, color=NAVY)
    header_row(ws, start, ["Kategori", "Toplam", *STATUSES[1:], "Bekliyor"])
    for i, name in enumerate(data["kategoriler"].values(), start=start + 1):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=f'=COUNTIF({rng_cat},A{i})')
        for j, s in enumerate(STATUSES[1:] + ["Bekliyor"], start=3):
            ws.cell(row=i, column=j, value=f'=COUNTIFS({rng_cat},$A{i},{rng_status},"{s}")')
        for c in range(1, 8):
            cell = ws.cell(row=i, column=c)
            cell.font = font()
            cell.border = BORDER
            if c > 1:
                cell.alignment = Alignment(horizontal="center")
    end = start + len(data["kategoriler"])
    tot = end + 1
    ws.cell(row=tot, column=1, value="Toplam").font = font(bold=True)
    for c in range(2, 8):
        col = get_column_letter(c)
        cell = ws.cell(row=tot, column=c, value=f"=SUM({col}{start + 1}:{col}{end})")
        cell.font = font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER
    ws.cell(row=tot, column=1).border = BORDER

    legend = tot + 2
    ws.cell(row=legend, column=1, value="Nasıl kullanılır?").font = font(size=12, bold=True, color=NAVY)
    notes = [
        "1. 'İhtiyaçlar' sayfasında yalnızca Durum (H), Kanıt (J) ve Not (K) sütunları güncellenir.",
        "2. Durum açılır listeden seçilir; ✓ sütunu ve bu sayfadaki sayılar formülle otomatik hesaplanır.",
        "3. Kalıcı güncelleme için aynı değişiklik ihtiyaclar.yaml dosyasına da yazılmalıdır.",
        "4. 'Bekliyor' = üzerinde çalışılıyor ya da henüz doğrulanamadı; kaldırılan ve birleştirilen ihtiyaçlar 'Değişiklik Geçmişi' sayfasındadır.",
        "5. Cümle yapısı için 'EARS Kalıpları' sayfasına bakın.",
    ]
    for k, note in enumerate(notes, start=legend + 1):
        ws.cell(row=k, column=1, value=note).font = font(color="404040")


def build_ears_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("EARS Kalıpları")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 80
    header_row(ws, 1, ["Tip", "Kalıp", "Örnek"])
    rows = [
        ("Her zaman", "<Özne>, <nesne> <eylem>-malıdır/-melidir.", "Web arayüzü, her cevabın altında yalnızca o cevapta atıf yapılan kaynakları numaralı olarak listelemelidir."),
        ("Olay", "<Olay> -dığında/-diğinde, <özne> <eylem>-malıdır.", "Corpus'ta bir belge eklendiğinde, değiştiğinde veya silindiğinde, sistem yalnızca o belgenin dense ve BM25 indeks kayıtlarını güncellemelidir."),
        ("Durum", "<Durum> iken, <özne> <eylem>-malıdır.", "Cihaz ayarı 'auto' iken, sistem modelleri uyumlu bir GPU varsa GPU'da, yoksa CPU'da çalıştırmalıdır."),
        ("İstenmeyen", "<İstenmeyen durum> durumunda, <özne> <eylem>-malıdır.", "Soru TÜBİTAK destek programları ve Ar-Ge mevzuatı kapsamı dışında olduğu durumunda, sistem soruyu kapsam dışı olduğunu belirterek reddetmelidir."),
        ("Opsiyonel", "<Özellik> etkinleştirildiğinde, <özne> <eylem>-malıdır.", "Kök bulma seçeneği etkinleştirildiğinde, sistem BM25 belirteçlerine Türkçe kök bulma uygulamalıdır."),
    ]
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = font(bold=(c == 1))
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=8, column=1, value="Kurallar").font = font(bold=True, color=NAVY)
    for r, rule in enumerate(
        [
            "Her ihtiyaç tek bir yükümlülük içerir (tek '-malıdır/-melidir').",
            "Özne açıkça yazılır: Sistem, Corpus, Test seti, Web arayüzü, REST API, Komut satırı arayüzü, Değerlendirme modülü, Cevap değerlendirmesi, Rapor üretici, Docker paketi, Barındırılan sürüm…",
            "İhtiyaç aracı kullanan kişinin (soru soran, RAG öğrenen, aracı çalıştıran, katkıda bulunan) bakışıyla yazılır; geliştirme süreci ihtiyaç değildir.",
            "İhtiyaç 'ne' yapılacağını söyler; 'nasıl' yapılacağı kabul kriterinde ve tasarımda yer alır.",
            "Her ihtiyaç bir doğrulama yöntemiyle (Test, Gösterim, İnceleme, Analiz) ve somut bir kanıtla eşleşir.",
        ],
        start=9,
    ):
        ws.cell(row=r, column=1, value=f"• {rule}").font = font()


def build_glossary_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Sözlük")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    header_row(ws, 1, ["Terim", "Açıklama"])
    terms = [
        ("RAG", "Retrieval-Augmented Generation: dil modelinin cevabı, aranan belge parçalarına dayanarak üretmesi."),
        ("Parça (chunk)", "Belgenin arama ve embedding için bölündüğü metin birimi."),
        ("Embedding", "Metnin anlamını temsil eden sayısal vektör."),
        ("Dense arama", "Sorgu ve parça vektörleri arasındaki benzerliğe dayalı anlamsal arama."),
        ("BM25", "Kelime eşleşmesi ve kelime nadirliğine dayalı klasik (sparse) arama algoritması."),
        ("Hibrit arama", "Dense ve BM25 sonuçlarının birleştirilmesi."),
        ("RRF", "Reciprocal Rank Fusion: sıralamaları 1/(k+sıra) ile birleştiren yöntem (k=60)."),
        ("Yeniden sıralayıcı", "Aday parçaları soru ile birlikte okuyup yeniden puanlayan cross-encoder model."),
        ("Guardrail", "Uydurma cevabı ve kapsam dışı kullanımı engelleyen güvenlik önlemi."),
        ("Kanıt alıntısı", "Test setinde doğru cevabın geçtiği birebir metin parçası; arama isabetini ölçmek için kullanılır."),
        ("Recall@k", "Doğru kanıtın ilk k sonuç içinde bulunduğu soruların oranı."),
        ("MRR", "Mean Reciprocal Rank: doğru kanıtın ilk göründüğü sıranın tersinin ortalaması."),
        ("Hakem modeli", "Cevabın beklenen cevapla uyumunu ve kaynağa sadakatini puanlayan dil modeli (LLM-as-judge)."),
        ("Sadakat", "Cevaptaki iddiaların atıf yapılan kaynaklarda gerçekten geçme oranı (faithfulness)."),
        ("Doğru ret / yanlış ret", "Reddedilmesi gereken soruyu reddetme oranı / cevaplanabilir bir soruyu gereksiz yere reddetme oranı."),
        ("Hazır indeks paketi", "`rag export-index` çıktısı: koleksiyon anlık görüntüleri, BM25 indeksleri ve SHA-256 manifesti; ilk açılışta ingest yerine yüklenir."),
        ("EARS", "Easy Approach to Requirements Syntax: ihtiyaç cümleleri için standart kalıp ailesi."),
        ("Artımlı indeksleme", "Yalnızca yeni, değişen veya silinen belgelerin indekste güncellenmesi."),
    ]
    for r, (term, desc) in enumerate(terms, start=2):
        ws.cell(row=r, column=1, value=term).font = font(bold=True)
        ws.cell(row=r, column=2, value=desc).font = font()
        for c in (1, 2):
            ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")


def build_history_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Değişiklik Geçmişi")
    for col, width in zip("ABCD", (12, 10, 20, 90)):
        ws.column_dimensions[col].width = width
    header_row(ws, 1, ["Tarih", "Sürüm", "Yazan", "Değişiklik"])
    for r, row in enumerate(data.get("gecmis", []), start=2):
        for c, key in enumerate(("tarih", "surum", "yazan", "degisiklik"), start=1):
            cell = ws.cell(row=r, column=c, value=row[key])
            cell.font = font()
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_markdown(data: dict) -> None:
    cats = data["kategoriler"]
    lines = [
        MD_START,
        "",
        "> Bu bölüm `docs/requirements/build.py` tarafından `ihtiyaclar.yaml` dosyasından üretilir; elle düzenlemeyin.",
        "> Takip için Excel: [`docs/ihtiyaclar.xlsx`](ihtiyaclar.xlsx).",
        "",
    ]
    current = None
    for req in data["ihtiyaclar"]:
        if req["kat"] != current:
            current = req["kat"]
            lines += ["", f"### {cats[current]}", "", "| ID | Tip | İhtiyaç | Öncelik | Doğrulama | Durum |", "|---|---|---|---|---|---|"]
        lines.append(
            f"| {req['id']} | {req['tip']} | {req['ihtiyac']} | {req['oncelik']} | {req['dogrulama']} | {req['durum']} |"
        )
    lines += ["", MD_END]
    text = MD_DOC.read_text(encoding="utf-8")
    pattern = re.compile(re.escape(MD_START) + r".*?" + re.escape(MD_END), re.S)
    if not pattern.search(text):
        raise ValueError("Markers not found in ihtiyaclar.md")
    MD_DOC.write_text(pattern.sub(lambda _: "\n".join(lines), text), encoding="utf-8", newline="\n")


def main() -> None:
    data = load()
    wb = Workbook()
    last = build_requirements_sheet(wb, data)
    build_summary_sheet(wb, data, last)
    build_ears_sheet(wb)
    build_glossary_sheet(wb)
    build_history_sheet(wb, data)
    wb.move_sheet("Özet", offset=-wb.index(wb["Özet"]))
    wb.active = 0
    wb.save(XLSX_OUT)
    build_markdown(data)
    print(f"{len(data['ihtiyaclar'])} requirements -> {XLSX_OUT.name}, {MD_DOC.name}")


if __name__ == "__main__":
    main()
