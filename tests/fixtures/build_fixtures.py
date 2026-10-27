"""Build the binary test fixtures (small synthetic PDFs and an XLSX).

Run once after changing this file and commit the outputs::

    .venv\\Scripts\\python.exe tests/fixtures/build_fixtures.py

The PDFs use the core Helvetica font, which only covers Latin-1, so their text is ASCII.
Real corpus files are never committed; integration tests read them from ``data/raw``.
"""

from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent

METIN_PAGES = [
    [
        "TUBITAK",
        "UYGULAMA ESASLARI",
        "Amac",
        "MADDE 1 - (1) Bu esaslarin amaci, ornek destek programina iliskin usul ve esaslari",
        "belirlemektir. Program kapsaminda projelere hibe destek saglanir.",
        "1/2",
    ],
    [
        "Kapsam",
        "MADDE 2 - (1) Bu esaslar, programa basvuran kuruluslarin projelerinin",
        "degerlendirilmesi, izlenmesi ve sonuclandirilmasina iliskin kurallari kapsar.",
        "2/2",
    ],
]

KILAVUZ_PAGES = [
    [
        "1. Giris",
        "Bu kilavuz, ornek programa basvuru yapacak kuruluslar icin hazirlanmistir.",
        "Basvuru oncesinde kilavuzun tamami okunmalidir.",
        "2. Basvuru Kosullari",
        "Turkiye'de yerlesik sermaye sirketleri basvuru yapabilir.",
        "Basvurular cevrim ici sistem uzerinden yapilir.",
    ],
    [
        "3. Destek Orani",
        "Destek orani en fazla yuzde yetmis bestir.",
        "Destek suresi en fazla on sekiz aydir.",
    ],
]


def _text_pdf(pages: list[list[str]], path: Path) -> None:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for lines in pages:
        pdf.add_page()
        for line in lines:
            pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def _empty_pdf(path: Path) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.rect(20, 20, 100, 60)
    pdf.output(str(path))


def _overprinted_pdf(path: Path) -> None:
    """Fake bold made by printing the same line twice, as in some old corpus PDFs."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for _ in range(2):
        pdf.set_xy(10, 20)
        pdf.cell(0, 8, "Madde 1- Bu esaslar programin kurallarini belirler.")
    pdf.output(str(path))


def _limits_xlsx(path: Path) -> None:
    wb = Workbook()
    limits = wb.active
    assert limits is not None
    limits.title = "Limitler"
    limits.append(["TEYDEB DESTEK PROGRAMLARI"])
    limits.append([None, 1501, 1507])
    limits.append(["Proje Destek Süresi", "36 AY", "18 AY"])
    limits.append(["Destek Oranı", "KOBİ:%75\nBÜYÜK:%60", "KOBİ:%75"])
    program = wb.create_sheet("1501")
    program.append(["Proje Süresi Üst Sınırı", "Destek Oranı"])
    program.append(["36 AY", "KOBİ:%75"])
    wb.save(path)


def main() -> None:
    """Write every generated fixture next to this file."""
    (HERE / "pdf").mkdir(exist_ok=True)
    _text_pdf(METIN_PAGES, HERE / "pdf" / "metin.pdf")
    _text_pdf(KILAVUZ_PAGES, HERE / "pdf" / "kilavuz.pdf")
    _empty_pdf(HERE / "pdf" / "bos.pdf")
    _overprinted_pdf(HERE / "pdf" / "kalin.pdf")
    _limits_xlsx(HERE / "ust_limit.xlsx")


if __name__ == "__main__":
    main()
